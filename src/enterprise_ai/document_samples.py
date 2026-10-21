"""Deterministic synthetic CRM/HCM/SCM source corpus. Never real company policy."""
import argparse
import datetime
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path

TOPICS = {
    'crm': [('报价审批','折扣超过10%需要销售主管审批。'), ('客户分级','年度模拟采购额超过50万元标为重点客户。'),
            ('合同归档','签署后两个工作日内归档到合同台账。'), ('退换货','收到申请后三个工作日内记录核查结果。'),
            ('商机交接','交接记录应包含联系人、下一步动作和责任人。')],
    'hcm': [('入职清单','入职协调员确认工位和培训安排。'), ('休假申请','模拟休假由直属负责人审核后记录。'),
            ('费用报销','模拟报销需附用途和票据编号。'), ('培训登记','课程结束后记录学时和反馈。'),
            ('岗位变更','岗位变更需重新确认系统访问角色。')],
    'scm': [('补货规则','库存低于模拟安全线20件时发起补货建议。'), ('供应商准入','采购员记录资质核查和交付能力评估。'),
            ('到货验收','仓库按采购单核对数量和外观。'), ('采购审批','模拟采购金额超过3万元由采购主管审批。'),
            ('异常交付','发现延期后一个工作日内登记风险和替代安排。')],
}


def _pdf(title, paragraphs):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen.canvas import Canvas
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    out = io.BytesIO()
    canvas = Canvas(out, pagesize=letter, invariant=1)
    canvas.setTitle(title)
    y = 730
    for index, paragraph in enumerate([title, *paragraphs]):
        canvas.setFont('STSong-Light', 18 if index == 0 else 12)
        for start in range(0, len(paragraph), 36):
            canvas.drawString(54, y, paragraph[start:start+36]); y -= 22
        y -= 12
    canvas.save()
    return out.getvalue()


def _docx(title, paragraphs):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(12 if name == 'Normal' else 20)
        style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.add_heading(title, 0)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.core_properties.title = title
    doc.core_properties.author = 'Enterprise AI teaching corpus'
    fixed = datetime.datetime(2026,1,1, tzinfo=datetime.timezone.utc)
    doc.core_properties.created = doc.core_properties.modified = fixed
    out = io.BytesIO(); doc.save(out)
    normalized = io.BytesIO()
    with zipfile.ZipFile(out) as source, zipfile.ZipFile(normalized,'w',compression=zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2026,1,1,0,0,0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return normalized.getvalue()


def generate_corpus(output: Path, count=30, formats='mixed') -> dict:
    if type(count) is not int or count not in (30,300,1000) or formats not in ('md','mixed'):
        raise ValueError('count must be 30/300/1000; formats must be md/mixed')
    if formats == 'mixed':
        try:
            import docx, reportlab  # noqa: F401; fail before creating output
        except ImportError as exc:
            raise ValueError('install requirements/m08-documents.txt') from exc
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    documents = []
    for i in range(count):
        tenant = 1 + (i//15)%2
        domain = ('crm','hcm','scm')[(i%15)//5]
        topic, rule = TOPICS[domain][i%5]
        doc_id = f'{domain}-{i+1:04d}'
        title = f'{topic} · 合成租户{tenant}'
        paragraphs = ['仅用于教学的合成规则，不是真实企业制度或法律意见。',
                      f'文档标识：{doc_id}；租户：{tenant}；版本：1。',
                      f'业务规则：{rule}',
                      f'执行步骤：责任人登记申请，检查租户{tenant}的业务台账，记录审核人与处理时间。',
                      '异常处理：资料缺失时退回补充；未经审核不得标记完成。',
                      '追溯要求：保留申请编号、来源文档版本和审核记录。']
        suffix = ('md','pdf','docx')[i%3] if formats == 'mixed' else 'md'
        relative = f'tenant-{tenant}/{doc_id}.{suffix}'
        path = output/relative
        path.parent.mkdir(exist_ok=True)
        raw = (('# '+title+'\n\n'+'\n\n'.join(paragraphs)+'\n').encode()
               if suffix == 'md' else (_pdf(title,paragraphs) if suffix == 'pdf' else _docx(title,paragraphs)))
        with path.open('xb') as stream:
            stream.write(raw)
        documents.append({'tenant_id':tenant,'document_id':doc_id,'version':1,'op':'upsert',
                          'source_path':relative,'title':title,'domain':domain,
                          'roles':[{'crm':'sales','hcm':'hr','scm':'procurement'}[domain]],
                          'effective_from':'2026-01-01','effective_to':None,
                          'sha256':hashlib.sha256(raw).hexdigest()})
    catalog = {'schema_version':'m08-catalog-v1','documents':documents}
    with (output/'catalog.json').open('x',encoding='utf-8',newline='\n') as stream:
        json.dump(catalog,stream,ensure_ascii=False,sort_keys=True,indent=2); stream.write('\n')
    return catalog


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command',required=True)
    generate = sub.add_parser('generate')
    generate.add_argument('--output',type=Path,required=True)
    generate.add_argument('--count',type=int,default=30,choices=(30,300,1000))
    generate.add_argument('--formats',choices=('md','mixed'),default='mixed')
    args = parser.parse_args(argv)
    try:
        catalog = generate_corpus(args.output,args.count,args.formats)
        print(json.dumps({'documents':len(catalog['documents']),'output':str(args.output)}))
        return 0
    except (ValueError,OSError) as exc:
        print(str(exc),file=sys.stderr); return 1


if __name__ == '__main__':
    raise SystemExit(main())
