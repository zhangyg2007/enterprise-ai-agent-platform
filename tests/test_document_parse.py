import io
import importlib.util
import unittest
import zipfile

from enterprise_ai.document_parse import parse_document


class ParseTests(unittest.TestCase):
    def test_markdown_sections_fences_and_normalization(self):
        source = '# 制度\r\n\r\nCafe\u0301。\n\n```text\n# 这不是标题\n\n保留\n```\n\n## 附录\n正文'
        blocks = parse_document(source.encode(), '.md')
        self.assertEqual([s['block'] for s in blocks], list(range(1,len(blocks)+1)))
        self.assertTrue(all(s['page'] is None for s in blocks))
        self.assertEqual(blocks[1]['text'], 'Café。')
        self.assertIn('# 这不是标题\n\n保留', blocks[2]['text'])
        self.assertEqual(blocks[2]['section'], '制度')
        self.assertEqual(blocks[-1]['section'], '附录')

    def test_invalid_empty_and_limits(self):
        for raw, suffix in ((b'', '.md'), (b' \n', '.md'), (b'\xff', '.md'),
                            (b'a\x00b','.md'), (b'a', '.html'), (b'x'* (10*1024*1024+1), '.md'),
                            (b'a'*500001, '.md'), (b'bad','.docx'), (b'bad','.pdf')):
            with self.subTest(suffix=suffix, size=len(raw)):
                with self.assertRaises(ValueError):
                    parse_document(raw, suffix)

    def test_zip_escape_duplicate_and_ratio(self):
        for names in (['../oops'], ['word/document.xml','word/document.xml'], ['word/vbaProject.bin']):
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer,'w') as archive:
                for name in names:
                    archive.writestr(name, 'x')
            with self.assertRaises(ValueError):
                parse_document(buffer.getvalue(), '.docx')
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer,'w',compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('word/document.xml', 'x'*100000)
        with self.assertRaises(ValueError):
            parse_document(buffer.getvalue(), '.docx')

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'python-docx optional')
    def test_docx_order_and_unsupported_content(self):
        from docx import Document
        from docx.oxml import OxmlElement
        doc = Document()
        doc.add_heading('规则', 1)
        doc.add_paragraph('开始')
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text, row.cells[1].text = '字段','数值'
        doc.add_paragraph('结束')
        out = io.BytesIO(); doc.save(out)
        blocks = parse_document(out.getvalue(), '.docx')
        self.assertEqual([b['text'] for b in blocks], ['规则','开始','字段 | 数值','结束'])
        doc.element.body.append(OxmlElement('w:ins'))
        out = io.BytesIO(); doc.save(out)
        with self.assertRaises(ValueError):
            parse_document(out.getvalue(), '.docx')

    @unittest.skipUnless(all(importlib.util.find_spec(n) for n in ('pypdf','reportlab')), 'PDF optional libraries')
    def test_pdf_pages_and_empty_page_rejected(self):
        from reportlab.pdfgen.canvas import Canvas
        out = io.BytesIO(); canvas = Canvas(out)
        for word in ('First', 'Second'):
            canvas.drawString(50,700,word); canvas.showPage()
        canvas.save()
        blocks = parse_document(out.getvalue(), '.pdf')
        self.assertEqual([b['page'] for b in blocks], [1,2])
        out = io.BytesIO(); canvas = Canvas(out); canvas.showPage(); canvas.save()
        with self.assertRaisesRegex(ValueError, 'empty|OCR'):
            parse_document(out.getvalue(), '.pdf')

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'python-docx optional')
    def test_docx_equation_must_not_be_silently_dropped(self):
        from docx import Document
        from docx.oxml import OxmlElement
        for equation_tag in ('m:oMath','m:oMathPara'):
            with self.subTest(tag=equation_tag):
                doc = Document()
                paragraph = doc.add_paragraph('Charge calculation: ')
                math = OxmlElement('m:oMath')
                run = OxmlElement('m:r'); text = OxmlElement('m:t')
                text.text = 'fee = principal * 0.15'
                run.append(text); math.append(run)
                if equation_tag=='m:oMathPara':
                    wrapper = OxmlElement(equation_tag); wrapper.append(math); math=wrapper
                paragraph._p.append(math)
                raw = io.BytesIO(); doc.save(raw)
                with self.assertRaisesRegex(ValueError,'unsupported DOCX'):
                    parse_document(raw.getvalue(),'.docx')


if __name__ == '__main__':
    unittest.main()
